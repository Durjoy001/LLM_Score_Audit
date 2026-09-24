#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Bootstrap 95% CIs for the Spearman correlations reported in RQ1 and RQ3.

Neither evaluate_cohens_kappa.py (RQ1) nor evaluate_kickstarter_correlation.py
(RQ3) bootstraps a CI for Spearman's rho -- only for weighted kappa / ICC.
This script adds that CI, reusing the exact same _bootstrap_ci() routine
(1000 resamples, percentile method, numpy Generator seeded 42) already used
for kappa/ICC in evaluate_cohens_kappa.py, applied to _spearman()'s rho.

It does NOT re-run either pipeline or touch their stored outputs -- it
reconstructs the exact pairs each pipeline scored (from the checked-in
evaluation_report.xlsx per_item sheet for RQ1, and from the same loader
functions evaluate_kickstarter_correlation.py uses for RQ3), verifies the
reconstructed rho matches the already-reported value, and only then adds
the bootstrap CI on top.

Output: spearman_bootstrap_cis.csv (repo root) with columns
  test_name, rho, ci_lower, ci_upper, n
"""
import csv
import json
import sys
from pathlib import Path

import openpyxl

BASE_DIR = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(BASE_DIR))
sys.path.insert(0, str(BASE_DIR / "src" / "tools"))

from evaluate_cohens_kappa import _spearman, _bootstrap_ci  # noqa: E402
import evaluate_kickstarter_correlation as ekc  # noqa: E402

DIMENSIONS = ["team", "objective", "strategy", "advantages", "feasibility"]
OUT_CSV = BASE_DIR / "spearman_bootstrap_cis.csv"


def rq1_pairs():
    """Reconstruct the exact 12 human/AI pairs per column from the checked-in
    RQ1 report (results/Dataset1/evaluation/evaluation_report.xlsx, per_item sheet)."""
    xlsx_path = BASE_DIR / "results" / "Dataset1" / "evaluation" / "evaluation_report.xlsx"
    summary_path = BASE_DIR / "results" / "Dataset1" / "evaluation" / "evaluation_summary.json"
    summary = json.loads(summary_path.read_text(encoding="utf-8"))

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["per_item"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col_idx = {name: i for i, name in enumerate(header)}

    pairs_by_col = {c: [] for c in DIMENSIONS + ["overall_ranking"]}
    for row in ws.iter_rows(min_row=2, values_only=True):
        for col in pairs_by_col:
            h = row[col_idx[f"human_{col}"]]
            a = row[col_idx[f"ai_{col}"]]
            if h in (None, "") or a in (None, ""):
                continue
            pairs_by_col[col].append((float(h), float(a)))

    # Sanity check: reconstructed pairs must reproduce the already-reported rho.
    for col in pairs_by_col:
        rho, _ = _spearman(pairs_by_col[col])
        reported = summary["metrics"][col]["spearman_r"]
        assert abs(rho - reported) < 1e-6, (
            f"RQ1 {col}: reconstructed rho={rho} != reported rho={reported}"
        )
    return pairs_by_col


def rq3_pairs():
    """Reconstruct the exact 141-proposal pairs per column the same way
    evaluate_kickstarter_correlation.py does, and verify against the
    checked-in results/Kickstarter/evaluation/evaluation_summary.json."""
    summary_path = BASE_DIR / "results" / "Kickstarter" / "evaluation" / "evaluation_summary.json"
    summary = json.loads(summary_path.read_text(encoding="utf-8"))["combined"]

    funded = ekc.load_funded_pct()
    ai_scores = ekc.load_ai_scores()
    all_pids = sorted(funded.keys())
    scored_pids = [p for p in all_pids if p in ai_scores]

    pairs_by_col = {c: [] for c in DIMENSIONS + ["overall_ranking"]}
    for pid in scored_pids:
        s = ai_scores[pid]
        pct = funded[pid]
        pairs_by_col["overall_ranking"].append((s["overall_ranking"], pct))
        for d in DIMENSIONS:
            pairs_by_col[d].append((s[d], pct))

    reported = {"overall_ranking": summary["overall_spearman"]}
    reported.update(summary["dimension_spearman"])
    for col in pairs_by_col:
        rho, _ = _spearman(pairs_by_col[col])
        rep_rho = reported[col]["rho"]
        assert abs(rho - rep_rho) < 1e-6, (
            f"RQ3 {col}: reconstructed rho={rho} != reported rho={rep_rho}"
        )
    return pairs_by_col


def main() -> None:
    rows = []
    for rq_label, pairs_fn in (("RQ1", rq1_pairs), ("RQ3", rq3_pairs)):
        pairs_by_col = pairs_fn()
        for col in DIMENSIONS + ["overall_ranking"]:
            pairs = pairs_by_col[col]
            rho, _ = _spearman(pairs)
            ci_lo, ci_hi = _bootstrap_ci(pairs, lambda p: _spearman(p)[0])
            label = "overall" if col == "overall_ranking" else col
            rows.append({
                "test_name": f"{rq_label}_{label}",
                "rho": rho,
                "ci_lower": ci_lo,
                "ci_upper": ci_hi,
                "n": len(pairs),
            })

    with OUT_CSV.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["test_name", "rho", "ci_lower", "ci_upper", "n"])
        writer.writeheader()
        writer.writerows(rows)

    print(f"[OK] wrote {OUT_CSV} ({len(rows)} rows)")
    for r in rows:
        print(r)


if __name__ == "__main__":
    main()
