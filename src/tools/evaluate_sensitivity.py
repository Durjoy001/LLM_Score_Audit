# -*- coding: utf-8 -*-
"""
Evaluation: Sensitivity Analysis

Tests whether the LLM's proposal evaluation responds logically to changes
in business proposal metrics.  Takes N processed proposals (baseline + variants
that have already been run through the full pipeline) and checks:

  1. Did the right DIMENSIONS change?
     (e.g. a TAM doubling should shift 'objective', not 'team')
  2. Did scores move in the RIGHT DIRECTION?
     (improvement in a metric → higher score for that dimension)
  3. Is the MAGNITUDE proportional?
     (small change → small shift; big change → bigger shift)
  4. Are there UNEXPECTED changes or MISSED changes?

Inputs:
  --pids           Comma-separated list of PIDs (must have reports + llm_scores)
  --baseline_pid   Which PID is the baseline (default: first in list)
  --model          LLM model for the assessment calls

Outputs:
  <out_dir>/sensitivity_summary.json
  <out_dir>/sensitivity_report.xlsx
"""

import json
import uuid
import argparse
import sys
from pathlib import Path
from datetime import datetime
from typing import Any, Dict, List, Optional

from dotenv import load_dotenv

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from src.backend.utils.llm_provider import active_provider, chat_json, default_model, provider_api_key
from src.backend.utils.score_extraction import extract_report_scores
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parents[2]
DATA_DIR = BASE_DIR / "src" / "data"
REPORT_DIR = DATA_DIR / "reports"
SCORES_DIR = DATA_DIR / "llm_scores"
PROPOSALS_DIR = DATA_DIR / "proposals"
EVAL_DIR = DATA_DIR / "evaluations"
RUNS_DIR = EVAL_DIR / "runs"

load_dotenv()
PROVIDER = active_provider()

ORDINAL_DIMS = ["team", "objective", "strategy", "advantages", "feasibility"]
ALL_DIMS = ORDINAL_DIMS + ["overall_ranking"]
ALL_FIELDS = ALL_DIMS + ["verdict"]
MAX_REPORT_CHARS = 16_000
# Minimum absolute delta to be treated as a real score change (filters noise)
MIN_DELTA_THRESHOLD = 0.01


def _log(level: str, msg: str) -> None:
    print(f"[{level}] {msg}", flush=True)


def now_str() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, obj: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")


def _truncate(text: str, max_chars: int) -> str:
    text = (text or "").strip()
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + "\n... [TRUNCATED] ..."


def _extract_evaluation_sections(report_md: str) -> str:
    """Return only Sections 0 and 1 of the report (before the Q&A evidence block).

    Scoping comparison to the executive summary and expert review avoids sending
    the verbose Q&A evidence to the LLM assessor, which is both costly and
    irrelevant for sensitivity comparison.
    """
    for marker in ("\n## 2.", "## 2."):
        idx = report_md.find(marker)
        if idx != -1:
            return report_md[:idx].strip()
    return report_md


# ---------------------------------------------------------------------------
# Data loading
# ---------------------------------------------------------------------------

def _load_proposal(pid: str) -> Dict[str, Any]:
    report_path = REPORT_DIR / f"{pid}_final_report.md"
    score_path = SCORES_DIR / pid / "llm_scores.json"
    meta_path = PROPOSALS_DIR / f"{pid}.meta.json"

    if not report_path.exists():
        raise FileNotFoundError(f"No final report for pid={pid}: {report_path}")

    meta: Dict[str, Any] = {}
    if meta_path.exists():
        try:
            meta = read_json(meta_path)
        except Exception:
            pass

    report_text = report_path.read_text(encoding="utf-8")

    # Primary: llm_scores.json (written by --use_llm mode, integer scores)
    # Fallback: extract 0-1 floats directly from the report
    if score_path.exists():
        score_data = read_json(score_path)
        source_filename = (
            score_data.get("meta", {}).get("source_filename")
            or meta.get("original_filename")
            or pid
        )
        scores = score_data.get("scores", {})
    else:
        source_filename = meta.get("original_filename") or pid
        scores = extract_report_scores(report_text)

    return {
        "pid": pid,
        "label": source_filename,
        "scores": scores,
        "report_excerpt": _truncate(_extract_evaluation_sections(report_text), MAX_REPORT_CHARS),
    }


# ---------------------------------------------------------------------------
# Score deltas
# ---------------------------------------------------------------------------

def _compute_deltas(baseline: Dict, variant: Dict) -> Dict[str, Any]:
    """Return per-field comparison dicts: variant score vs baseline."""
    deltas: Dict[str, Any] = {}

    for dim in ORDINAL_DIMS:
        bv = baseline.get(dim)
        vv = variant.get(dim)
        if bv is None or vv is None:
            continue
        try:
            bv, vv = float(bv), float(vv)
        except (TypeError, ValueError):
            continue
        d = round(vv - bv, 4)
        deltas[dim] = {
            "baseline": round(bv, 4),
            "variant": round(vv, 4),
            "delta": d,
            "direction": "up" if d > 0 else "down" if d < 0 else "unchanged",
        }

    br = baseline.get("overall_ranking")
    vr = variant.get("overall_ranking")
    if br is not None and vr is not None:
        try:
            br, vr = float(br), float(vr)
            d = round(vr - br, 4)
            deltas["overall_ranking"] = {
                "baseline": round(br, 4),
                "variant": round(vr, 4),
                "delta": d,
                "direction": "up" if d > 0 else "down" if d < 0 else "unchanged",
            }
        except (TypeError, ValueError):
            pass

    bverd = str(baseline.get("verdict") or "").strip().upper()
    vverd = str(variant.get("verdict") or "").strip().upper()
    if bverd and vverd:
        deltas["verdict"] = {
            "baseline": bverd,
            "variant": vverd,
            "changed": bverd != vverd,
            "label": f"{bverd}→{vverd}" if bverd != vverd else "unchanged",
        }

    return deltas


def _fmt_deltas(deltas: Dict) -> str:
    lines = []
    for field, info in deltas.items():
        if field == "verdict":
            lines.append(f"  verdict: {info['label']}")
        else:
            d = info.get("delta", 0)
            sign = "+" if d >= 0 else ""
            arrow = "↑" if info["direction"] == "up" else "↓" if info["direction"] == "down" else "→"
            lines.append(
                f"  {field}: {info['baseline']} → {info['variant']}  "
                f"(delta {sign}{d} {arrow})"
            )
    return "\n".join(lines) if lines else "  (no numeric changes detected)"


# ---------------------------------------------------------------------------
# LLM assessment
# ---------------------------------------------------------------------------

def _llm_assess(
    baseline: Dict,
    variant: Dict,
    deltas: Dict,
    provider: str,
    model: str,
) -> Dict[str, Any]:
    if not provider_api_key(provider):
        return {
            "sensitivity_quality": "unknown",
            "direction_assessment": "unknown",
            "appropriately_sensitive": None,
            "dimensions_correctly_affected": [],
            "unexpected_changes": [],
            "missed_changes": [],
            "key_observations": [],
            "overall_verdict": f"Skipped — no {provider.upper()} API key configured.",
        }

    changed_dims = [
        k for k, v in deltas.items()
        if k not in ("verdict",) and abs(v.get("delta", 0)) >= MIN_DELTA_THRESHOLD
    ]
    unchanged_dims = [
        k for k, v in deltas.items()
        if k not in ("verdict",) and abs(v.get("delta", 0)) < MIN_DELTA_THRESHOLD
    ]

    system_prompt = (
        "You are an expert evaluator assessing whether an AI proposal-analysis system "
        "is appropriately sensitive to changes in business proposal content and metrics.\n\n"
        "You will receive:\n"
        "  1. Score changes (variant minus baseline) across rubric dimensions\n"
        "  2. A baseline report excerpt\n"
        "  3. A variant report excerpt\n\n"
        "Your job: judge whether the score shifts are logically coherent with the "
        "differences visible in the two reports — as a domain expert would expect.\n\n"
        "Return ONLY valid JSON, no markdown fences."
    )

    user_prompt = (
        "=== SCORE CHANGES (variant − baseline) ===\n"
        f"{_fmt_deltas(deltas)}\n\n"
        f"Dimensions that changed: {changed_dims or 'none'}\n"
        f"Dimensions unchanged:    {unchanged_dims or 'none'}\n\n"
        "=== BASELINE REPORT EXCERPT ===\n"
        f"{baseline['report_excerpt']}\n\n"
        "=== VARIANT REPORT EXCERPT ===\n"
        f"{variant['report_excerpt']}\n\n"
        "Based on the OBSERVABLE DIFFERENCES between the two reports, assess:\n"
        "1. Are score changes directionally correct? (metric improves → score rises?)\n"
        "2. Did the RIGHT dimensions change? (market metric → objectives/strategy, "
        "not team)\n"
        "3. Did any dimension change for NO visible reason in the reports?\n"
        "4. Did any dimension FAIL to change when it clearly should have?\n"
        "5. Is magnitude proportional to the scale of the difference?\n\n"
        "Return exactly this JSON:\n"
        "{\n"
        '  "sensitivity_quality": "good" | "partial" | "poor",\n'
        '  "direction_assessment": "correct" | "mixed" | "incorrect",\n'
        '  "appropriately_sensitive": true | false,\n'
        '  "dimensions_correctly_affected": ["dim", ...],\n'
        '  "unexpected_changes": ["dim: reason", ...],\n'
        '  "missed_changes": ["dim: reason", ...],\n'
        '  "key_observations": ["...", ...],\n'
        '  "overall_verdict": "one-sentence summary"\n'
        "}"
    )

    try:
        result = chat_json(
            provider=provider,
            model=model,
            temperature=0.1,
            max_tokens=800,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
        )
        # Normalise required keys
        for key in ("unexpected_changes", "missed_changes", "key_observations",
                    "dimensions_correctly_affected"):
            if not isinstance(result.get(key), list):
                result[key] = []
        return result
    except Exception as exc:
        _log("WARN", f"LLM assessment failed: {exc}")
        return {
            "sensitivity_quality": "unknown",
            "direction_assessment": "unknown",
            "appropriately_sensitive": None,
            "dimensions_correctly_affected": [],
            "unexpected_changes": [],
            "missed_changes": [],
            "key_observations": [],
            "overall_verdict": f"Assessment failed: {exc}",
            "error": str(exc),
        }


# ---------------------------------------------------------------------------
# Synthesis
# ---------------------------------------------------------------------------

def _synthesize(assessments: List[Dict]) -> Dict[str, Any]:
    if not assessments:
        return {"overall_sensitivity": "unknown", "summary": "No variants assessed."}

    quality_map = {"good": 2, "partial": 1, "poor": 0, "unknown": 1}
    avg = sum(quality_map.get(a.get("sensitivity_quality", "unknown"), 1)
              for a in assessments) / len(assessments)
    overall = "good" if avg >= 1.7 else "partial" if avg >= 0.8 else "poor"

    n_sensitive = sum(1 for a in assessments if a.get("appropriately_sensitive") is True)
    n_total = len(assessments)

    all_issues: List[str] = []
    for a in assessments:
        all_issues.extend(a.get("unexpected_changes", []))
        all_issues.extend(a.get("missed_changes", []))
    # deduplicate preserving order
    seen: set = set()
    unique_issues = [x for x in all_issues if not (x in seen or seen.add(x))]

    return {
        "overall_sensitivity": overall,
        "variants_assessed": n_total,
        "variants_sensitive": n_sensitive,
        "sensitivity_rate": round(n_sensitive / n_total, 4) if n_total else 0.0,
        "all_issues": unique_issues,
        "summary": (
            f"{n_sensitive}/{n_total} variants show appropriate sensitivity. "
            f"Overall quality: {overall}."
        ),
    }


# ---------------------------------------------------------------------------
# Excel output
# ---------------------------------------------------------------------------

_HEADER_FILL = PatternFill("solid", fgColor="1E293B")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_DELTA_UP_FILL = PatternFill("solid", fgColor="DCFCE7")
_DELTA_DOWN_FILL = PatternFill("solid", fgColor="FEE2E2")


def _header_row(ws, values: List[str]) -> None:
    ws.append(values)
    for cell in ws[ws.max_row]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def write_excel(
    out_path: Path,
    baseline: Dict,
    variants: List[Dict],
    deltas_list: List[Dict],
    assessments: List[Dict],
    overall: Dict,
) -> None:
    wb = Workbook()

    # ── Sheet 1: score comparison ────────────────────────────────────────
    ws = wb.active
    ws.title = "score_comparison"
    variant_labels = [f"V{i+1} — {v['label'][:40]}" for i, v in enumerate(variants)]
    _header_row(ws, ["Dimension", f"Baseline — {baseline['label'][:40]}"] + variant_labels)

    for field in ALL_FIELDS:
        row = [field, baseline["scores"].get(field, "—")]
        for v in variants:
            row.append(v["scores"].get(field, "—"))
        ws.append(row)

    ws.column_dimensions["A"].width = 18
    for col_idx in range(2, 2 + len(variants) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 28

    # ── Sheet 2: score deltas ────────────────────────────────────────────
    ws2 = wb.create_sheet("score_deltas")
    _header_row(ws2, ["Dimension"] + [f"V{i+1} delta vs baseline" for i in range(len(variants))])

    for field in ALL_FIELDS:
        row: List[Any] = [field]
        for idx, dmap in enumerate(deltas_list):
            info = dmap.get(field, {})
            if field == "verdict":
                row.append(info.get("label", "—"))
            else:
                d = info.get("delta")
                if d is None:
                    row.append("—")
                else:
                    sign = "+" if d >= 0 else ""
                    row.append(f"{sign}{d}  ({info.get('direction', '?')})")
        ws2.append(row)

    # Colour-code deltas
    for row_idx in range(2, ws2.max_row + 1):
        field = ws2.cell(row_idx, 1).value
        if field == "verdict":
            continue
        for col_idx, dmap in enumerate(deltas_list, start=2):
            info = dmap.get(field, {})
            d = info.get("delta")
            if d is None:
                continue
            cell = ws2.cell(row_idx, col_idx)
            if d > 0:
                cell.fill = _DELTA_UP_FILL
            elif d < 0:
                cell.fill = _DELTA_DOWN_FILL

    ws2.column_dimensions["A"].width = 18
    for col_idx in range(2, 2 + len(variants) + 1):
        ws2.column_dimensions[get_column_letter(col_idx)].width = 24

    # ── Sheet 3: LLM assessments ─────────────────────────────────────────
    ws3 = wb.create_sheet("llm_assessments")
    _header_row(ws3, [
        "variant", "file", "sensitivity_quality", "direction_assessment",
        "appropriately_sensitive", "dimensions_correctly_affected",
        "unexpected_changes", "missed_changes",
        "key_observations", "overall_verdict",
    ])
    for i, (v, a) in enumerate(zip(variants, assessments), start=1):
        ws3.append([
            f"V{i}",
            v["label"],
            a.get("sensitivity_quality", "—"),
            a.get("direction_assessment", "—"),
            str(a.get("appropriately_sensitive", "—")),
            "; ".join(a.get("dimensions_correctly_affected", [])),
            "; ".join(a.get("unexpected_changes", [])),
            "; ".join(a.get("missed_changes", [])),
            "; ".join(a.get("key_observations", [])),
            a.get("overall_verdict", "—"),
        ])
    for col in ws3.columns:
        ws3.column_dimensions[col[0].column_letter].width = 30

    # ── Sheet 4: overall summary ─────────────────────────────────────────
    ws4 = wb.create_sheet("summary")
    _header_row(ws4, ["metric", "value"])
    for k, v in [
        ("baseline", baseline["label"]),
        ("n_variants", len(variants)),
        ("overall_sensitivity", overall["overall_sensitivity"]),
        ("variants_sensitive", f"{overall['variants_sensitive']}/{overall['variants_assessed']}"),
        ("sensitivity_rate", overall["sensitivity_rate"]),
        ("summary", overall["summary"]),
    ]:
        ws4.append([k, v])
    ws4.append([])
    ws4.append(["issues"])
    for issue in overall.get("all_issues", []):
        ws4.append(["", issue])
    ws4.column_dimensions["A"].width = 22
    ws4.column_dimensions["B"].width = 60

    wb.save(out_path)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> None:
    ap = argparse.ArgumentParser(description="Sensitivity analysis for LLM proposal evaluation.")
    ap.add_argument("--pids", required=True,
                    help="Comma-separated list of PIDs (baseline first if --baseline_pid omitted).")
    ap.add_argument("--baseline_pid", default="",
                    help="PID of the baseline proposal (default: first in --pids).")
    ap.add_argument("--model", default="",
                    help="LLM model for assessment calls.")
    ap.add_argument("--provider", default="", help="LLM provider for assessment calls, e.g. openai or gemini.")
    ap.add_argument("--out_dir", default="",
                    help="Output directory (auto-generated if omitted).")
    args = ap.parse_args()
    provider = args.provider.strip().lower() or PROVIDER
    model = args.model or default_model(provider)

    pid_list = [p.strip() for p in args.pids.split(",") if p.strip()]
    if len(pid_list) < 2:
        raise RuntimeError("Sensitivity analysis requires at least 2 PIDs.")

    baseline_pid = args.baseline_pid.strip() or pid_list[0]
    if baseline_pid not in pid_list:
        raise RuntimeError(f"baseline_pid={baseline_pid!r} not in --pids list.")

    run_id = datetime.now().strftime("%Y%m%d_%H%M%S") + "_" + uuid.uuid4().hex[:6]
    out_dir = Path(args.out_dir) if args.out_dir else (RUNS_DIR / run_id)
    out_dir.mkdir(parents=True, exist_ok=True)

    _log("INFO", f"Loading {len(pid_list)} proposals...")
    all_data: Dict[str, Dict] = {}
    for pid in pid_list:
        _log("INFO", f"  Loading {pid}")
        all_data[pid] = _load_proposal(pid)

    baseline = all_data[baseline_pid]
    variants = [all_data[p] for p in pid_list if p != baseline_pid]

    _log("INFO", f"Baseline: {baseline['label']}")
    for i, v in enumerate(variants, 1):
        _log("INFO", f"  Variant {i}: {v['label']}")

    deltas_list: List[Dict] = []
    assessments: List[Dict] = []

    for i, variant in enumerate(variants, 1):
        _log("INFO", f"[{i}/{len(variants)}] Computing deltas for {variant['label']}...")
        deltas = _compute_deltas(baseline["scores"], variant["scores"])
        deltas_list.append(deltas)

        _log("INFO", f"[{i}/{len(variants)}] LLM assessment...")
        assessment = _llm_assess(baseline, variant, deltas, provider, model)
        assessments.append(assessment)
        _log("INFO", (
            f"  quality={assessment.get('sensitivity_quality')}  "
            f"direction={assessment.get('direction_assessment')}  "
            f"sensitive={assessment.get('appropriately_sensitive')}"
        ))

    overall = _synthesize(assessments)
    _log("INFO", f"Overall: {overall['summary']}")

    payload = {
        "meta": {
            "run_id": run_id,
            "generated_at": now_str(),
            "provider": provider,
            "model": model,
            "baseline_pid": baseline_pid,
            "baseline_label": baseline["label"],
            "n_variants": len(variants),
        },
        "overall": overall,
        "baseline": {
            "pid": baseline_pid,
            "label": baseline["label"],
            "scores": baseline["scores"],
        },
        "variants": [
            {
                "pid": v["pid"],
                "label": v["label"],
                "scores": v["scores"],
                "deltas": deltas_list[i],
                "assessment": assessments[i],
            }
            for i, v in enumerate(variants)
        ],
    }

    summary_path = out_dir / "sensitivity_summary.json"
    report_path = out_dir / "sensitivity_report.xlsx"
    write_json(summary_path, payload)
    write_excel(report_path, baseline, variants, deltas_list, assessments, overall)
    _log("OK", f"summary  → {summary_path}")
    _log("OK", f"report   → {report_path}")


if __name__ == "__main__":
    main()
