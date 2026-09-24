#!/usr/bin/env python3
"""
Workbook for checking the PLUS/MINUS edits by hand: for every proposal x sub-rubric, what text was
ADDED (PLUS; the manually corrected sentence where the original was misaligned), and what text is
REMOVED and what it is REPLACED WITH (MINUS; from the manual review in removed_text_review.py).

Read-only over saved results; makes no LLM calls.

Output: perturbation_review.xlsx
  Pass rates                 corrected rescoring per sub-rubric and totals, with the original rates alongside
  Strategy ... Feasibility   one sheet per dimension, one row per proposal x sub-rubric
"""

import importlib.util
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
DATA = ROOT / "src" / "data"
CONTENT_DIR = DATA / "sensitivity_results" / "adaptive_content"
RESULTS_JSON = DATA / "sensitivity_results" / "sensitivity_all_dimensions.json"
CORRECTED_JSON = DATA / "sensitivity_results" / "corrected_v2" / "results.json"
SINGLE_JSON = DATA / "sensitivity_results" / "corrected_v2" / "minus_single.json"
PLUS_SINGLE_JSON = DATA / "sensitivity_results" / "corrected_v2" / "plus_single.json"
OUT_XLSX = HERE / "perturbation_review.xlsx"

import sys  # noqa: E402
sys.path.insert(0, str(HERE))
from added_text_review import REVIEW  # noqa: E402  manual alignment review + corrected text
from removed_text_review import MINUS_REMOVE, MINUS_REPLACEMENT  # noqa: E402  manual MINUS review
from aggregate_subrubric_passrates import tally  # noqa: E402  same pass rule as the published figures

DIMENSIONS = ["strategy", "objectives", "advantages", "team", "feasibility"]
PIDS = ["p1", "p2", "p3", "p4", "p5", "p6", "p7", "p8", "pA", "pB", "pC", "pD"]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
PLUS_FILL = PatternFill("solid", fgColor="E2EFDA")
MINUS_FILL = PatternFill("solid", fgColor="FCE4D6")
BAND_FILL = PatternFill("solid", fgColor="F2F2F2")
NOTHING_TO_REMOVE = "(nothing to remove: the report has no passage supporting this sub-rubric)"
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


def clean(value):
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


def style_header(ws, n_cols):
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"


def main() -> int:
    runner_spec = importlib.util.spec_from_file_location("runner", ROOT / "run_sensitivity_all_dimensions.py")
    runner = importlib.util.module_from_spec(runner_spec)
    runner_spec.loader.exec_module(runner)

    wb = Workbook()

    # ── Pass rates ───────────────────────────────────────────────────────────
    # Corrected rescoring (rescore_corrected.py), with the original run's rates for comparison.
    ws = wb.active
    ws.title = "Pass rates"
    original = json.loads(RESULTS_JSON.read_text())["results"]
    checks = json.loads(CORRECTED_JSON.read_text())["checks"]
    headers = ["dimension", "sub_rubric", "label", "type",
               "PLUS passed", "PLUS tested", "PLUS rate", "PLUS excluded (baseline=5)",
               "MINUS passed", "MINUS tested", "MINUS rate", "MINUS excluded (baseline=1)",
               "MINUS untestable (nothing to remove)",
               "MINUS one-at-a-time passed", "MINUS one-at-a-time tested", "MINUS one-at-a-time rate",
               "PLUS one-at-a-time passed", "PLUS one-at-a-time tested", "PLUS one-at-a-time rate",
               "original PLUS rate", "original MINUS rate"]
    single_data = json.loads(SINGLE_JSON.read_text())
    single = {(r["pid"], r["sub_rubric"]): r["result"] for r in single_data["checks"]}
    plus_single_data = json.loads(PLUS_SINGLE_JSON.read_text())
    plus_single = {(r["pid"], r["sub_rubric"]): r["result"] for r in plus_single_data["checks"]}
    ws.append(headers)

    def summarise(rows):
        out = []
        for col in ("plus_result", "minus_result"):
            ok = sum(r[col] == "PASS" for r in rows)
            n = sum(r[col] in ("PASS", "FAIL") for r in rows)
            out += [ok, n, round(ok / n, 3) if n else "n/a", sum(r[col] == "EXCLUDED" for r in rows)]
        out.append(sum(r["minus_result"] == "UNTESTABLE" for r in rows))
        for lookup in (single, plus_single):
            results = [lookup.get((r["pid"], r["sub_rubric"])) for r in rows]
            ok = results.count("PASS")
            n = ok + results.count("FAIL")
            out += [ok, n, round(ok / n, 3) if n else "n/a"]
        return out

    def original_rates(pairs):
        out = []
        for direction in ("PLUS", "MINUS"):
            ok = n = 0
            for dim, key in pairs:
                a, b, _ = tally(original[dim], key, direction)
                ok, n = ok + a, n + b
            out.append(round(ok / n, 3) if n else "n/a")
        return out

    for dim in DIMENSIONS:
        for key, label, rtype in runner.DIMENSIONS[dim]["rubrics"]:
            rows = [r for r in checks if r["sub_rubric"] == key]
            ws.append([dim, key, label, rtype, *summarise(rows), *original_rates([(dim, key)])])

    ws.append([])
    ws.append(["TOTALS"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    keys = {d: [(d, k) for k, _, _ in runner.DIMENSIONS[d]["rubrics"]] for d in DIMENSIONS}
    types = {k: ty for d in DIMENSIONS for k, _, ty in runner.DIMENSIONS[d]["rubrics"]}
    for dim in DIMENSIONS + ["ALL"]:
        for rtype in ("all types", "doc-verifiable", "judgment-dep"):
            rows = [r for r in checks if dim in ("ALL", r["dimension"])
                    and rtype in ("all types", r["rubric_type"])]
            pairs = [pk for d in (DIMENSIONS if dim == "ALL" else [dim]) for pk in keys[d]
                     if rtype in ("all types", types[pk[1]])]
            ws.append([dim, "", "", rtype, *summarise(rows), *original_rates(pairs)])

    same = sum(r["baseline"] == r["baseline_repeat"] for r in checks)
    ws.append([])
    for line in (
        "Corrected rescoring: added and removed text as shown on the dimension sheets; PLUS text inserted before",
        "'##### Team and governance' for every dimension; baseline and variants scored identically (temperature 0,",
        "no expert-opinion JSON, neutral report IDs). Rates count PASS / (PASS + FAIL).",
        "MINUS untestable = the report had no passage supporting the sub-rubric, so nothing was removed.",
        f"Baseline scored twice: {same} of {len(checks)} sub-rubric scores were identical across the two runs.",
        "MINUS one-at-a-time: a separate report per sub-rubric with only that sub-rubric's passages removed "
        "(rescore_minus_single.py).",
        f"  In those reports the OTHER sub-rubrics dropped in {100 * single_data['spillover_drop_rate']:.0f}% of cases; "
        f"scores dropped between two identical baseline runs in {100 * single_data['noise_drop_rate']:.0f}% of cases.",
        "PLUS one-at-a-time: a separate report per sub-rubric with only that sub-rubric's sentence added "
        "(rescore_plus_single.py).",
        f"  In those reports the OTHER sub-rubrics rose in {100 * plus_single_data['spillover_rise_rate']:.0f}% of cases; "
        f"scores rose between two identical baseline runs in {100 * plus_single_data['noise_rise_rate']:.0f}% of cases.",
        "Original rates: the first run (run_sensitivity_all_dimensions.py --adaptive), for comparison.",
    ):
        ws.append([line])
    style_header(ws, len(headers))
    for col in (7, 11, 16, 19, 20, 21):
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=col).number_format = "0.0%"
    for i, w in enumerate([12, 36, 24, 15, 8, 8, 8, 12, 8, 8, 8, 12, 14, 10, 10, 10, 10, 10, 10, 10, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ── One sheet per dimension ──────────────────────────────────────────────
    headers = ["proposal", "sub_rubric", "type",
               "ADDED text (aimed at this sub-rubric)", "REMOVED text (original)", "REPLACED WITH"]
    for dim in DIMENSIONS:
        ws = wb.create_sheet(dim.capitalize())
        ws.append(headers)
        row_no = 1
        for band, pid in enumerate(PIDS):
            content = json.loads((CONTENT_DIR / f"{pid}_{dim}.json").read_text())
            coverage = content["plus"].get("rubric_coverage") or {}
            for key, label, rtype in runner.DIMENSIONS[dim]["rubrics"]:
                removed = MINUS_REMOVE.get(pid, {}).get(dim, {}).get(key, [])
                # Corrected sentence from the manual review where the original was misaligned.
                review = REVIEW.get(dim, {}).get(key, {}).get(pid)
                ws.append([clean(v) for v in [
                    pid, label, rtype,
                    review[1] if review else coverage.get(key, ""),
                    "\n---\n".join(removed) if removed else NOTHING_TO_REMOVE,
                    MINUS_REPLACEMENT[key] if removed else ""]])
                row_no += 1
                for col in range(1, len(headers) + 1):
                    c = ws.cell(row=row_no, column=col)
                    c.alignment = WRAP if col >= 4 else TOP
                    if band % 2 and col <= 3:
                        c.fill = BAND_FILL
                ws.cell(row=row_no, column=4).fill = PLUS_FILL
                for col in (5, 6):
                    ws.cell(row=row_no, column=col).fill = MINUS_FILL
        style_header(ws, len(headers))
        ws.auto_filter.ref = ws.dimensions
        for i, w in enumerate([9, 24, 15, 60, 60, 45], 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(OUT_XLSX)
    print(f"Wrote {OUT_XLSX}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
