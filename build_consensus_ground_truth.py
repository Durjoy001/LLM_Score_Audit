#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Task 1 — Build a 3-rater consensus ground truth to replace the original
single-rater ExpertsEvaluations.xlsx.

Reads the three rater workbooks in
DeltaProjectDatasets/ExpertsEvaluation/ExpertEvaluation{1,2,3}.xlsx, aligns them
by Proposal ID (1-8 = biotech proposals, 9-12 = semiconductor A/B/C/D), and:

  - median across the 3 raters for the five 1-5 dimension scores
  - mean across the 3 raters for the 1-20 overall ranking (Overall /20)
  - majority vote for Funded Y/N, with a flag whenever it is not unanimous

Outputs (all no-LLM, purely from cached expert Excel files):
  - ExpertsEvaluations_consensus.xlsx   (repo root; drop-in replacement that
        mirrors the original human_scores.xlsx schema AND file names so it
        matches the cached AI reports in evaluate_cohens_kappa.py)
  - resultsNew/Dataset1/ExpertsEvaluations_consensus.xlsx  (copy alongside results)
  - resultsNew/Dataset1/raters_long.csv       (raw long-format 3-rater table)
  - resultsNew/Dataset1/consensus_build.json  (consensus values + unanimity flags)
"""
import csv
import json
import statistics
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent
RATER_DIR = BASE / "DeltaProjectDatasets" / "ExpertsEvaluation"
TEMPLATE = BASE / "src" / "data" / "evaluations" / "human" / "human_scores.xlsx"
OUT_XLSX = BASE / "ExpertsEvaluations_consensus.xlsx"
RESULTS_DIR = BASE / "resultsNew" / "Dataset1"

DIMS = ["team", "objective", "strategy", "advantages", "feasibility"]

# Proposal ID in the rater files -> pid used by calibration_sweep.py, and the
# row order in the original human_scores.xlsx template.
#   rater IDs 1..8  -> biotech p1..p8
#   rater IDs 9..12 -> semiconductor A,B,C,D  (pA..pD)
RATER_ID_TO_PID = {
    1: "p1", 2: "p2", 3: "p3", 4: "p4", 5: "p5", 6: "p6", 7: "p7", 8: "p8",
    9: "pA", 10: "pB", 11: "pC", 12: "pD",
}


def _to_int(v):
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return None


def _verdict(v):
    s = str(v or "").strip().upper()
    return "Y" if s.startswith("Y") else ("N" if s.startswith("N") else "")


def load_rater(path: Path, rater_id: int):
    """Return {proposal_id:int -> row dict} for one rater workbook.

    Column layout (all three raters share it):
      0 Proposal ID | 1 File Name | 2 Project Name | 3-7 Team..Feasibility /5
      8 Raw Total /25 | 9 Overall /20 | 10 Funded Y/N
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    out = {}
    for r in rows[1:]:
        if r is None or r[0] is None:
            continue
        pid_num = _to_int(r[0])
        if pid_num is None or pid_num not in RATER_ID_TO_PID:
            continue
        out[pid_num] = {
            "rater_id": rater_id,
            "proposal_id": pid_num,
            "pid": RATER_ID_TO_PID[pid_num],
            "file_name": str(r[1] or "").strip(),
            "project_name": str(r[2] or "").strip(),
            "team": _to_int(r[3]),
            "objective": _to_int(r[4]),
            "strategy": _to_int(r[5]),
            "advantages": _to_int(r[6]),
            "feasibility": _to_int(r[7]),
            "raw_total_25": _to_int(r[8]),
            "overall_20": _to_int(r[9]),
            "verdict": _verdict(r[10]),
        }
    return out


def main():
    raters = {
        1: load_rater(RATER_DIR / "ExpertEvaluation1.xlsx", 1),
        2: load_rater(RATER_DIR / "ExpertEvaluation2.xlsx", 2),
        3: load_rater(RATER_DIR / "ExpertEvaluation3.xlsx", 3),
    }

    RESULTS_DIR.mkdir(parents=True, exist_ok=True)

    # ---- raw long-format table (rater_id, proposal_id, dimension, score) ----
    long_rows = []
    for rid, table in raters.items():
        for pid_num in sorted(table):
            rec = table[pid_num]
            for dim in DIMS:
                long_rows.append({
                    "rater_id": rid, "proposal_id": pid_num, "pid": rec["pid"],
                    "dimension": dim, "score": rec[dim],
                })
            long_rows.append({
                "rater_id": rid, "proposal_id": pid_num, "pid": rec["pid"],
                "dimension": "overall_ranking", "score": rec["overall_20"],
            })
    long_csv = RESULTS_DIR / "raters_long.csv"
    with long_csv.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["rater_id", "proposal_id", "pid", "dimension", "score"])
        w.writeheader()
        w.writerows(long_rows)

    # ---- consensus per proposal ----
    consensus = {}
    for pid_num in sorted(RATER_ID_TO_PID):
        recs = [raters[rid][pid_num] for rid in (1, 2, 3) if pid_num in raters[rid]]
        pid = RATER_ID_TO_PID[pid_num]
        dim_consensus = {d: int(statistics.median([r[d] for r in recs])) for d in DIMS}
        overall_vals = [r["overall_20"] for r in recs]
        overall_mean = round(statistics.mean(overall_vals), 4)
        # Raw Total /25 is, by definition, the sum of the five dimension scores,
        # so the consensus Raw Total = sum of the consensus (median) dimensions
        # (not the median of raters' raw totals, which need not equal it).
        raw_total = sum(dim_consensus[d] for d in DIMS)
        verdicts = [r["verdict"] for r in recs]
        n_yes = verdicts.count("Y")
        majority = "Y" if n_yes >= 2 else "N"
        unanimous = len(set(verdicts)) == 1
        consensus[pid] = {
            "proposal_id": pid_num,
            "pid": pid,
            **dim_consensus,
            "overall_20_mean": overall_mean,
            "raw_total_25": raw_total,
            "verdict_majority": majority,
            "verdict_by_rater": verdicts,
            "verdict_unanimous": unanimous,
            "verdict_flag": None if unanimous else f"NON-UNANIMOUS ({n_yes}Y/{3 - n_yes}N)",
        }

    flagged = [p for p, c in consensus.items() if not c["verdict_unanimous"]]

    # ---- drop-in xlsx: clone the template, overwrite score cells ----
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb.active
    # Template row map (1-indexed openpyxl rows) -> pid
    #   header 'DataSet 1' r1, header r2, data r3..r10 = p1..p8
    #   blank r11, 'DataSet 2' r12, data r13..r16 = pA..pD
    row_to_pid = {}
    for i, pid in enumerate(["p1", "p2", "p3", "p4", "p5", "p6", "p7", "p8"]):
        row_to_pid[3 + i] = pid
    for i, pid in enumerate(["pA", "pB", "pC", "pD"]):
        row_to_pid[13 + i] = pid
    # Column indices (1-indexed): D..H dims, I overall ranking, J funded, K raw total
    COL = {"team": 4, "objective": 5, "strategy": 6, "advantages": 7,
           "feasibility": 8, "overall": 9, "funded": 10, "raw": 11}
    for row_idx, pid in row_to_pid.items():
        c = consensus[pid]
        ws.cell(row=row_idx, column=COL["team"], value=c["team"])
        ws.cell(row=row_idx, column=COL["objective"], value=c["objective"])
        ws.cell(row=row_idx, column=COL["strategy"], value=c["strategy"])
        ws.cell(row=row_idx, column=COL["advantages"], value=c["advantages"])
        ws.cell(row=row_idx, column=COL["feasibility"], value=c["feasibility"])
        ws.cell(row=row_idx, column=COL["overall"], value=c["overall_20_mean"])
        ws.cell(row=row_idx, column=COL["funded"], value=c["verdict_majority"])
        ws.cell(row=row_idx, column=COL["raw"], value=c["raw_total_25"])
    wb.save(OUT_XLSX)
    wb.save(RESULTS_DIR / "ExpertsEvaluations_consensus.xlsx")

    (RESULTS_DIR / "consensus_build.json").write_text(
        json.dumps({
            "method": {
                "dimensions": "median across 3 raters (1-5)",
                "overall_ranking": "mean across 3 raters (Overall /20)",
                "verdict": "majority vote (>=2 of 3); flagged when not unanimous",
            },
            "id_mapping": RATER_ID_TO_PID,
            "consensus": consensus,
            "non_unanimous_verdict_proposals": flagged,
        }, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print(f"[OK] consensus xlsx -> {OUT_XLSX}")
    print(f"[OK] copy           -> {RESULTS_DIR / 'ExpertsEvaluations_consensus.xlsx'}")
    print(f"[OK] long-format    -> {long_csv}")
    print(f"[OK] build json     -> {RESULTS_DIR / 'consensus_build.json'}")
    print(f"\nNon-unanimous Funded Y/N: {flagged or 'none'}")
    print("\npid  team obj str adv fea  overall  verdict  (raters)")
    for pid in ["p1", "p2", "p3", "p4", "p5", "p6", "p7", "p8", "pA", "pB", "pC", "pD"]:
        c = consensus[pid]
        flag = "" if c["verdict_unanimous"] else "  <-- FLAG " + str(c["verdict_by_rater"])
        print(f"{pid:3s}   {c['team']}   {c['objective']}   {c['strategy']}   "
              f"{c['advantages']}   {c['feasibility']}    {c['overall_20_mean']:5.2f}    "
              f"{c['verdict_majority']}{flag}")


if __name__ == "__main__":
    main()
